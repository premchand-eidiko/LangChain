from __future__ import annotations

import csv
from pathlib import Path
from typing import List

from langchain_community.document_loaders import (
    CSVLoader,
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document as LangChainDocument
from pptx import Presentation


class PowerPointLoader:
    def __init__(self, path: str):
        self.path = path

    def load(self) -> List[LangChainDocument]:
        presentation = Presentation(self.path)
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text = "\n".join(
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            )
            if text:
                slides.append(
                    LangChainDocument(
                        page_content=text,
                        metadata={"slide": slide_number},
                    )
                )
        return slides


LOADERS = {
    ".pdf": PyPDFLoader,
    ".docx": Docx2txtLoader,
    ".pptx": PowerPointLoader,
    ".txt": TextLoader,
    ".csv": CSVLoader,
}


def load_document(path: str) -> List[LangChainDocument]:
    file_path = Path(path)
    loader_class = LOADERS.get(file_path.suffix.lower())
    if loader_class is None:
        raise ValueError("Unsupported document type")
    loader = loader_class(str(file_path))
    documents = loader.load()
    if file_path.suffix.lower() == ".csv" and not documents:
        with file_path.open(newline="", encoding="utf-8-sig") as csv_file:
            rows = [", ".join(row) for row in csv.reader(csv_file) if any(row)]
        documents = [LangChainDocument(page_content="\n".join(rows))] if rows else []
    if not any(document.page_content.strip() for document in documents):
        raise ValueError("The document contains no readable text")
    return documents
